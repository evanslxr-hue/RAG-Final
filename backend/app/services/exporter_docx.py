from docx import Document


def write_docx(path: str, title: str, sections: list[tuple[str, str]]) -> None:
    doc = Document()
    doc.add_heading(title, level=1)
    for heading, body in sections:
        doc.add_heading(heading, level=2)
        doc.add_paragraph(body)
    doc.save(path)
